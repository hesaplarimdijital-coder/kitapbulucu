"""Write the sample files that LibreOffice produces, so the reader is tested against a real
second writer and not only against openpyxl and hand-made XML.

Usage: make_office_fixtures.py OUT_DIR   (normally app/src/test/resources/fixtures)

Needs `soffice` on the PATH. Every book here is invented.

- ornek-libreoffice.xlsx: saved by LibreOffice Calc (its own shared strings, styles and
  date formats; formulas carry cached results).
- ornek-libreoffice.xls: the old binary format, which the app must refuse with a clear message.
- ornek-libreoffice.csv: Calc's default text export (comma, quotes, UTF-8).
"""
import datetime
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

ROWS = [
    ("Kitap Adı", "Yazar", "Raf", "Adet", "Yıl", "Tarih", "Not"),
    ("Islıklı Kuşların Sokağı", "Nevin Çağlar", "A1", 1, "2011", datetime.date(1998, 3, 15), "Çatı katı; ikinci baskı"),
    ("Dağ \"Sessizlik\"", "Ferit Ağaoğlu", "B2", 2, "1953", datetime.datetime(1953, 4, 1), ""),
    ("Hâlâ Bekleyen Mektup", "Şükrü Ilgaz", "C3", 1, "", "", "İmzalı"),
    ("Toplam", "", "", "=SUM(D2:D4)", "", "", ""),
]


def source_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kitaplar"
    for row in ROWS:
        ws.append(list(row))
    ws["F2"].number_format = "DD.MM.YYYY"
    ws["F3"].number_format = "MMM/YY"
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "@"
    wb.create_sheet("Notlar")["B2"] = "İkinci sayfa"
    wb.save(path)


def convert(src, out_dir, target):
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    try:
        subprocess.run(
            ["soffice", "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", target, "--outdir", out_dir, src],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def main(out_dir):
    work = tempfile.mkdtemp(prefix="office-fixtures-")
    try:
        src = os.path.join(work, "kaynak.xlsx")
        source_workbook(src)
        converted = os.path.join(work, "out") # never the source folder: Calc would skip the xlsx
        os.mkdir(converted)
        for target, ext in (("xlsx:Calc MS Excel 2007 XML", "xlsx"), ("xls:MS Excel 97", "xls"), ("csv:Text - txt - csv (StarCalc):44,34,76,1", "csv")):
            convert(src, converted, target)
            shutil.copyfile(os.path.join(converted, f"kaynak.{ext}"), os.path.join(out_dir, f"ornek-libreoffice.{ext}"))
            print("wrote", f"ornek-libreoffice.{ext}")
    finally:
        shutil.rmtree(work, ignore_errors=True)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "app/src/test/resources/fixtures")
