"""Write the made-up sample workbooks used by the app's tests.

Usage: make_fixtures.py OUT_DIR   (normally app/src/test/resources/fixtures)

Every book here is invented. The real lists never enter this public repository.

- ornek-kutuphane.xlsx: the app's own layout (header row, nine columns).
- ornek-eski-bicim.xlsx: the old hand-kept layout: no header, two columns, capitals,
  a duplicated row, "İMZALI" inside the title, a cell Excel turned into a date, a header row
  lost in the middle, stray spaces.
"""
import datetime
import os
import sys

import openpyxl
from openpyxl.styles import Font

COLUMNS = ["Kitap Adı", "Yazar", "Raf", "Adet", "İmzalı", "Yayınevi", "Yıl", "ISBN", "Not"]

BOOKS = [
    # title, author, shelf, copies, signed, publisher, year, isbn, note
    ("Islıklı Kuşların Sokağı", "Nevin Çağlar", "A1", 1, "", "Mavi Defter", "2011", "9786050000017", ""),
    ("Sefil Fenerciler", "Halit Oruç", "A1", 2, "", "", "", "", ""),
    ("Hâlâ Bekleyen Mektup", "Şükrü Ilgaz", "A2", 1, "Evet", "Kumsal", "1998", "", "Yazarın kendi el yazısıyla"),
    ("İğde Ağacının Gölgesi", "Gülşen Ünlü", "B3", 1, "", "", "", "", ""),
    ("Kürk Yakalı Saatçi", "Ömer Işıkçı", "B3", 1, "", "Kuyu & Kalem", "2005", "9789750000025", ""),
    ("Çatıdaki Rasathane", "Ayşe İnce", "C2", 3, "", "", "", "", ""),
    ("Dağ Başındaki <Sessizlik>", "Ferit Ağaoğlu", "", 1, "", "", "", "", "Kapak yırtık"),
    ("Öğle Uykusu Üzerine Denemeler", "Işıl Öztürk", "3-4", 1, "", "", "", "", ""),
    ("Yağmurun İstanbul'u", "Cemal Kılıç", "D4", 1, "Evet", "", "", "", "Kitapçı Murat'a imzalı"),
    ("=Formül Sanılan Kitap", "Deniz Yalın", "", 1, "", "", "", "", ""),
    ("Uzun Bir Adı Olan ve Satır Sonuna Kadar Giden, Bitmek Bilmeyen Kitabın Adı", "Rüya Eşsiz", "E1", 1, "", "", "", "", ""),
    ("Kâtip Çelebi'nin Defteri", "", "E1", 1, "", "", "1950", "", "Yazarı bilinmiyor"),
]

OLD_LAYOUT = [
    ("ISLIKLI KUŞLARIN SOKAĞI", "NEVİN ÇAĞLAR"),
    ("SEFİL FENERCİLER ", "HALİT ORUÇ"),
    ("SEFİL FENERCİLER", "HALİT ORUÇ"),
    ("HÂLÂ BEKLEYEN MEKTUP - İMZALI", "ŞÜKRÜ ILGAZ"),
    ("TARLA DERGİSİ SAYI 1", datetime.datetime(1953, 4, 1)),
    ("KİTAP", "YAZAR"),
    ("ÇATIDAKİ  RASATHANE", "AYŞE İNCE"),
    ("YAĞMURUN İSTANBUL'U - CEMAL BEY'E İMZALI", "CEMAL KILIÇ"),
    ("YAZARSIZ BİR ALMANAK", None),
]


def make_library(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Babam"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for book in BOOKS:
        sheet.append(list(book))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column != 4:
                cell.number_format = "@"
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # keep "=Formül…" a text, not a formula
    sheet.freeze_panes = "A2"
    workbook.save(path)


def make_old_layout(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sayfa1"
    for title, author in OLD_LAYOUT:
        sheet.append([title, author])
    sheet["B5"].number_format = "mmm/yy"
    workbook.create_sheet("Sayfa2")
    workbook.save(path)


def main(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    make_library(os.path.join(out_dir, "ornek-kutuphane.xlsx"))
    make_old_layout(os.path.join(out_dir, "ornek-eski-bicim.xlsx"))
    print("written to", out_dir)


if __name__ == "__main__":
    main(sys.argv[1])
